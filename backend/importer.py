"""Local Excel/ZIP import with 1C-code matching and explicit data quality."""

from __future__ import annotations

import re
import math
from hashlib import sha256
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import UploadFile
from openpyxl import load_workbook

from .replenishment import CalculationInput, Inbound, Operation, Stockout
from .storage import save_preview
from .import_layouts import find_header, flat_views, normalize, TableView


MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6, "июл": 7, "авг": 8, "сен": 9, "сент": 9, "окт": 10, "ноя": 11, "дек": 12}


def _code(value: Any) -> str | None:
    if value is None:
        return None
    code = str(value).strip()
    return code or None


def _same_warehouse(value: Any, warehouse_scope: str) -> bool:
    """Compare user input and 1C warehouse values without case sensitivity."""
    return str(value or "").strip().casefold() == warehouse_scope.strip().casefold()


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    try:
        number = float(str(value).replace(" ", "").replace("\xa0", "").replace(",", "."))
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    match = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(value))
    if match:
        try:
            return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _month(value: Any) -> str | None:
    text = str(value or "").lower()
    year = re.search(r"20\d{2}", text)
    if not year:
        return None
    for prefix, month in MONTHS.items():
        if text.startswith(prefix):
            return f"{year.group()}-{month:02d}"
    return None


def _warning(code: str, severity: str, message: str, sku: str | None = None) -> dict:
    return {"code": code, "severity": severity, "message": message, "sku": sku}


def _detect(name: str, sheet: Any) -> str:
    preview_rows = list(
        sheet.iter_rows(
            min_row=1,
            max_row=min(30, sheet.max_row),
            values_only=True,
        )
    )
    first = [str(x or "").lower().strip() for x in (preview_rows[0] if preview_rows else ())]
    joined = " ".join(first)
    preview_text = " ".join(
        str(value or "").lower().strip()
        for row in preview_rows
        for value in row
    )
    month_columns = sum(_month(value) is not None for value in first)
    if (
        "норм. коэф." in preview_text and "месяц" in preview_text
    ):
        return "seasonality"
    if "количество" in joined and "склад" in joined and "документ" in joined:
        return "sales_detail"
    if "мин. разр." in joined:
        return "moq"
    if "поступление до" in joined:
        return "inbound"
    if (
        month_columns >= 6
        and "номенклатура.код" in first
        and any(value in {"ед.", "ед", "единица"} for value in first)
    ):
        return "monthly_stock"
    if month_columns >= 6 and "номенклатура.код" in first:
        return "monthly_sales"
    if "начало" in joined and "конец" in joined:
        return "stockouts"
    if "категор" in joined:
        return "categories"
    if "остат" in joined and ("дата" in joined or "на дату" in joined):
        return "current_stock"
    return "unknown"


def _parse_sheet(name: str, sheet: Any, detected: str, warehouse_scope: str, data: dict, issues: list[dict], counts: Counter) -> None:
    source = {"file": name, "sheet": sheet.title}
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    if detected == "sales_detail":
        seen: set[tuple] = set()
        skipped: Counter = Counter()
        skipped_examples: dict[str, list[int]] = defaultdict(list)
        customer_index = next((index for index, value in enumerate(header) if str(value or "").lower().strip() == "customer_id" or "обезлич" in str(value or "").lower()), None)
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not row or all(x is None for x in row):
                continue
            code = _code(row[3] if len(row) > 3 else None)
            sale_date = _date(row[0])
            quantity = _number(row[7] if len(row) > 7 else None)
            warehouse = str(row[6] or "").strip() if len(row) > 6 else ""
            document = str(row[2] or "") if len(row) > 2 else ""
            if not code or not sale_date or quantity is None:
                skipped["invalid"] += 1
                if len(skipped_examples["invalid"]) < 3:
                    skipped_examples["invalid"].append(row_no)
                continue
            if not _same_warehouse(warehouse, warehouse_scope):
                counts["other_warehouse_rows"] += 1
                continue
            if not document.startswith("Расходная накладная"):
                skipped["other_document"] += 1
                if len(skipped_examples["other_document"]) < 3:
                    skipped_examples["other_document"].append(row_no)
                continue
            if quantity <= 0:
                skipped["non_positive"] += 1
                if len(skipped_examples["non_positive"]) < 3:
                    skipped_examples["non_positive"].append(row_no)
                continue
            customer_id = _code(row[customer_index]) if customer_index is not None and len(row) > customer_index else None
            key = (sale_date, str(row[1]), document, code, quantity, warehouse, customer_id)
            if key in seen:
                issues.append(_warning("DUPLICATE_SALE_ROW", "error", f"{name}:{sheet.title}:{row_no}: операции с одинаковыми реквизитами; уточните, повтор это или разные строки документа. Расчёт по коду заблокирован.", code))
            seen.add(key)
            record = {"sku": code, "date": sale_date.isoformat(), "quantity": quantity, "name": str(row[4] or ""),
                      "unit": str(row[5] or ""), "warehouse": warehouse, "document": document, "document_number": str(row[1] or ""),
                      "customer_id": customer_id,
                      "source": {**source, "row": row_no}}
            data["sales_detail"].append(record)
            counts["sales_detail"] += 1
        if skipped["invalid"]:
            counts["invalid_sale_rows"] += skipped["invalid"]
            issues.append(_warning(
                "INVALID_SALE_ROWS",
                "warning",
                f"{name}:{sheet.title}: {skipped['invalid']} строк пропущено: нет кода, даты или количества (примеры строк: {', '.join(map(str, skipped_examples['invalid']))}).",
            ))
        if skipped["other_document"]:
            counts["excluded_non_sale_documents"] += skipped["other_document"]
            issues.append(_warning(
                "NON_SALE_DOCUMENTS_EXCLUDED",
                "warning",
                f"{name}:{sheet.title}: {skipped['other_document']} строк с типом документа, отличным от расходной накладной, не включены в обычный спрос (примеры строк: {', '.join(map(str, skipped_examples['other_document']))}).",
            ))
        if skipped["non_positive"]:
            counts["excluded_non_positive_sales"] += skipped["non_positive"]
            issues.append(_warning(
                "NON_POSITIVE_SALES_EXCLUDED",
                "warning",
                f"{name}:{sheet.title}: {skipped['non_positive']} строк с нулевым или отрицательным количеством не включены в обычный спрос; знак не меняется через abs() (примеры строк: {', '.join(map(str, skipped_examples['non_positive']))}).",
            ))
    elif detected in {"monthly_stock", "monthly_sales"}:
        code_col = 2 if detected == "monthly_stock" else 1
        first_data = 2
        month_columns = [(index, period) for index, value in enumerate(header) if (period := _month(value))]
        for row_no, row in enumerate(sheet.iter_rows(min_row=first_data, values_only=True), first_data):
            code = _code(row[code_col] if len(row) > code_col else None)
            if not code:
                continue
            unit = str(row[1] or "") if detected == "monthly_stock" else None
            for index, period in month_columns:
                value = _number(row[index] if len(row) > index else None)
                if value is None:
                    continue  # Missing is retained as absent, never converted to zero.
                data[detected].append({"sku": code, "month": period, "quantity": value, "unit": unit,
                                       "kind": "beginning_stock" if detected == "monthly_stock" else "monthly_sales",
                                       "source": {**source, "row": row_no, "column": index + 1}})
                counts[detected] += 1
    elif detected == "moq":
        invalid_moq_rows: list[int] = []
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            code = _code(row[1] if len(row) > 1 else None)
            if not code:
                continue
            minimum = _number(row[4] if len(row) > 4 else None)
            if minimum is None or minimum < 0:
                invalid_moq_rows.append(row_no)
                continue
            data["moq"].append({"sku": code, "minimum": minimum, "supplier_article": str(row[2] or "") if len(row) > 2 else "",
                                "source": {**source, "row": row_no}})
            counts["moq"] += 1
        if invalid_moq_rows:
            counts["invalid_moq_rows"] += len(invalid_moq_rows)
            issues.append(_warning(
                "INVALID_MOQ_ROWS",
                "warning",
                f"{name}:{sheet.title}: {len(invalid_moq_rows)} строк MOQ пропущено: минимальная партия пуста или отрицательна (примеры строк: {', '.join(map(str, invalid_moq_rows[:5]))}).",
            ))
    elif detected == "inbound":
        dates = [_date(value) for value in header]
        documents = [str(value or "").split("(")[0].strip() for value in header]
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            code = _code(row[0] if row else None)
            if not code:
                continue
            for index in range(3, len(header)):
                quantity = _number(row[index] if len(row) > index else None)
                if quantity is None or quantity == 0:
                    continue
                if not dates[index] or quantity < 0:
                    issues.append(_warning("INVALID_INBOUND", "warning", f"{name}:{sheet.title}:{row_no}: неизвестная дата или отрицательное количество", code))
                    continue
                note = str(row[2] or "") if len(row) > 2 else ""
                data["inbound"].append({"sku": code, "quantity": quantity, "expected_date": dates[index].isoformat(),
                                        "document": documents[index], "supplier_article": str(row[1] or ""),
                                        "unit_conversion_uncertain": "БУХТАМИ" in note.upper() and "МЕТРАЖОМ" in note.upper(),
                                        "source": {**source, "row": row_no, "column": index + 1}})
                counts["inbound"] += 1
    elif detected == "seasonality":
        layout = next(((n, row) for n, row in enumerate(sheet.iter_rows(max_row=min(30, sheet.max_row), values_only=True), 1)
                       if any(normalize(v) == "нормкоэф" for v in row)), None)
        if not layout:
            return
        header_no, season_header = layout
        month_col = next((i for i, v in enumerate(season_header) if normalize(v) == "месяц"), None)
        if month_col is None:
            issues.append(_warning("INVALID_SEASONALITY", "error", f"{name}:{sheet.title}: в строке нормированных коэффициентов нужна колонка «Месяц»."))
            return
        coef_col = next(i for i, v in enumerate(season_header) if normalize(v) == "нормкоэф")
        for row_no, row in enumerate(sheet.iter_rows(min_row=header_no + 1, values_only=True), header_no + 1):
            month_text = str(row[month_col] or "").strip().casefold()
            month = next((value for prefix, value in MONTHS.items() if month_text.startswith(prefix)), None)
            if month is None:
                continue
            coefficient = _number(row[coef_col] if len(row) > coef_col else None)
            if coefficient is None or coefficient <= 0:
                issues.append(_warning("INVALID_SEASONALITY", "error", f"{name}:{sheet.title}:{row_no}: нет нормированного коэффициента"))
                continue
            previous = data["seasonality"].get(str(month))
            if previous and previous["coefficient"] != coefficient:
                issues.append(_warning("CONFLICTING_SEASONALITY", "error", f"Разные коэффициенты сезонности для месяца {month}; требуется выбор источника."))
                continue
            data["seasonality"][str(month)] = {"coefficient": coefficient, "source": {**source, "row": row_no, "column": coef_col + 1}}
            counts["seasonality"] += 1
    elif detected == "current_stock":
        columns = {str(x or "").lower().strip(): index for index, x in enumerate(header)}
        code_index = next((v for k, v in columns.items() if "код" in k), None)
        quantity_index = next((v for k, v in columns.items() if "остат" in k), None)
        date_index = next((v for k, v in columns.items() if "дата" in k), None)
        warehouse_index = next((v for k, v in columns.items() if "склад" in k), None)
        if code_index is None or quantity_index is None or date_index is None:
            issues.append(_warning("INVALID_CURRENT_STOCK_FILE", "error", f"{name}:{sheet.title}: требуются код, остаток и дата"))
            return
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if warehouse_index is not None and not _same_warehouse(row[warehouse_index], warehouse_scope):
                counts["other_warehouse_stock_rows"] += 1
                continue
            code = _code(row[code_index] if len(row) > code_index else None)
            quantity = _number(row[quantity_index] if len(row) > quantity_index else None)
            snapshot_date = _date(row[date_index] if len(row) > date_index else None)
            if not code:
                continue
            if quantity is None or quantity < 0 or not snapshot_date:
                issues.append(_warning("INVALID_CURRENT_STOCK", "warning", f"{name}:{sheet.title}:{row_no}: неверный остаток или дата", code))
                continue
            unit_index = columns.get("ед.")
            data["current_stock"].append({"sku": code, "quantity": quantity, "date": snapshot_date.isoformat(), "warehouse": warehouse_scope,
                                          "unit": str(row[unit_index] or "").strip() if unit_index is not None else "", "source": {**source, "row": row_no}})
            counts["current_stock"] += 1
    elif detected == "categories":
        columns = {str(x or "").lower().strip(): index for index, x in enumerate(header)}
        code_index = next((v for k, v in columns.items() if "код" in k), None)
        category_index = next((v for k, v in columns.items() if "категор" in k), None)
        if code_index is None or category_index is None:
            return
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            code = _code(row[code_index] if len(row) > code_index else None)
            category = str(row[category_index] or "").strip() if len(row) > category_index else ""
            if code and category:
                data["categories"].append({"sku": code, "category": category, "source": {**source, "row": row_no}})
                counts["categories"] += 1
    elif detected == "stockouts":
        columns = {str(x or "").lower().strip(): index for index, x in enumerate(header)}
        code_index = next((v for k, v in columns.items() if "код" in k), None)
        start_index = next((v for k, v in columns.items() if "начал" in k or "start" in k), None)
        end_index = next((v for k, v in columns.items() if "конец" in k or "end" in k), None)
        if None in (code_index, start_index, end_index):
            issues.append(_warning("INVALID_STOCKOUT_FILE", "error", f"{name}:{sheet.title}: требуются код, начало и конец"))
            return
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if len(row) > 3 and row[3] and not _same_warehouse(row[3], warehouse_scope):
                continue
            code = _code(row[code_index] if len(row) > code_index else None)
            start = _date(row[start_index] if len(row) > start_index else None)
            end = _date(row[end_index] if len(row) > end_index else None)
            if not code:
                continue
            if not start or not end or end < start:
                issues.append(_warning("INVALID_STOCKOUT", "warning", f"{name}:{sheet.title}:{row_no}: неверный интервал", code))
                continue
            data["stockouts"].append({"sku": code, "start": start.isoformat(), "end": end.isoformat(), "confirmed": True, "source": {**source, "row": row_no}})
            counts["stockouts"] += 1
    elif detected == "inbound_rows":
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            code = _code(row[0])
            if not code or all(v is None for v in row[1:4]):
                continue
            if row[6] and not _same_warehouse(row[6], warehouse_scope):
                continue
            quantity, arrival_date = _number(row[1]), _date(row[2])
            if quantity is None or quantity < 0 or not arrival_date:
                issues.append(_warning("INVALID_INBOUND", "error", f"{name}:{sheet.title}:{row_no}: нужны количество и дата поставки.", code))
                continue
            data["inbound"].append({"sku": code, "quantity": quantity, "expected_date": arrival_date.isoformat(),
                                    "document": str(row[3] or ""), "supplier_article": str(row[4] or ""),
                                    "unit": str(row[5] or ""), "unit_conversion_uncertain": False,
                                    "source": {**source, "row": row_no}})
            counts["inbound"] += 1


def _merge_sources(data, issues, counts):
    sale_sources = {}
    overlap_skus = set()
    for record in data["sales_detail"]:
        key = tuple(record.get(f) for f in ("sku", "date", "document", "document_number", "quantity", "warehouse", "customer_id"))
        source = (record["source"]["file"], record["source"]["sheet"])
        if key in sale_sources and sale_sources[key] != source:
            overlap_skus.add(record["sku"])
        sale_sources[key] = source
    for sku in sorted(overlap_skus):
        issues.append(_warning("OVERLAPPING_SALES_SOURCES", "error", "Одна операция встречается в разных источниках; выберите одну детализацию продаж.", sku))
    # Repeated snapshots/settings in a combined sales table are facts, not additive quantities.
    for kind, fields in {
        "current_stock": ("sku", "warehouse", "date", "quantity", "unit"),
        "categories": ("sku", "category"), "moq": ("sku", "minimum", "supplier_article"),
        "stockouts": ("sku", "start", "end"),
        "monthly_sales": ("sku", "month", "quantity"),
        "monthly_stock": ("sku", "month", "quantity", "unit"),
    }.items():
        unique = {}
        for record in data[kind]:
            key = tuple(record.get(f) for f in fields)
            if key not in unique:
                unique[key] = record
            else:
                unique[key].setdefault("additional_sources", []).append(record["source"])
        removed = len(data[kind]) - len(unique)
        data[kind] = list(unique.values())
        if removed:
            counts[f"merged_repeated_{kind}"] = removed
        if kind in counts:
            counts[kind] = len(data[kind])
    for kind, keys, value in [("current_stock", ("sku", "warehouse", "date"), "quantity"),
                               ("categories", ("sku",), "category"),
                               ("monthly_sales", ("sku", "month"), "quantity"),
                               ("monthly_stock", ("sku", "month"), "quantity")]:
        variants = defaultdict(set)
        for record in data[kind]:
            variants[tuple(record.get(k) for k in keys)].add(record[value])
        for key, values in variants.items():
            if len(values) > 1:
                issues.append(_warning("SOURCE_VALUE_CONFLICT", "error", f"{kind}: противоречащие значения по коду {key[0]}; требуется уточнить источник.", key[0]))
    units = defaultdict(set)
    for kind in ("sales_detail", "current_stock", "inbound"):
        for record in data[kind]:
            if record.get("unit"):
                units[record["sku"]].add(normalize(record["unit"]))
    for sku, values in units.items():
        if len(values) > 1:
            issues.append(_warning("SOURCE_UNIT_CONFLICT", "error", "Разные единицы в источниках; требуется подтверждённое преобразование.", sku))


async def preview_files(files: list[UploadFile], warehouse_scope: str) -> dict:
    data: dict[str, Any] = {key: ([] if key != "seasonality" else {}) for key in
                            ("sales_detail", "monthly_sales", "monthly_stock", "inbound", "moq", "seasonality", "current_stock", "categories", "stockouts")}
    issues: list[dict] = []
    counts: Counter = Counter()
    file_info: list[dict] = []
    unsupported: set[str] = set()
    workbook_blobs: list[tuple[str, bytes]] = []
    for uploaded in files:
        blob = await uploaded.read()
        filename = uploaded.filename or "unnamed"
        if filename.rsplit("/", 1)[-1].startswith("~$"):
            issues.append(_warning(
                "TEMPORARY_EXCEL_FILE_SKIPPED",
                "info",
                f"{filename}: временный файл Excel пропущен.",
            ))
            continue
        if filename.lower().endswith(".zip"):
            try:
                archive = ZipFile(BytesIO(blob))
                workbook_blobs.extend((entry.filename.rsplit("/", 1)[-1], archive.read(entry)) for entry in archive.infolist() if entry.filename.lower().endswith(".xlsx") and not entry.filename.startswith("__MACOSX/"))
            except BadZipFile as error:
                raise ValueError(f"{filename}: unreadable ZIP") from error
        elif filename.lower().endswith(".xlsx"):
            workbook_blobs.append((filename, blob))
        else:
            raise ValueError(f"{filename}: expected .xlsx or .zip")
    if not workbook_blobs:
        raise ValueError("No Excel files found")
    unique_blobs: list[tuple[str, bytes]] = []
    seen_digests: set[str] = set()
    for filename, blob in workbook_blobs:
        if filename.startswith("~$"):
            continue
        digest = sha256(blob).hexdigest()
        if digest in seen_digests:
            issues.append(_warning("DUPLICATE_FILE_SKIPPED", "warning", f"{filename}: тот же файл уже есть в этой загрузке; повтор пропущен."))
            continue
        seen_digests.add(digest)
        unique_blobs.append((filename, blob))
    workbook_blobs = unique_blobs
    for filename, blob in workbook_blobs:
        try:
            workbook = load_workbook(BytesIO(blob), read_only=True, data_only=True)
        except Exception as error:
            issues.append(_warning("UNREADABLE_WORKBOOK", "error", f"{filename}: {error}"))
            file_info.append({"file_name": filename, "sheet_names": [], "detected_type": "unknown"})
            continue
        kinds: list[str] = []
        for sheet in workbook:
            header_info = find_header(sheet)
            views = []
            if header_info:
                header_no, raw, fields, duplicates = header_info
                if duplicates:
                    issues.append(_warning("AMBIGUOUS_COLUMNS", "error", f"{filename}:{sheet.title}: несколько колонок с одинаковым смыслом: {', '.join(sorted(duplicates))}."))
                    kinds.append("unknown")
                    continue
                month_indexes = [i for i, value in enumerate(raw) if _month(value)]
                if month_indexes:
                    # Monthly stock is recognized by the subheading, not just a unit column.
                    subheaders = " ".join(str(v or "").casefold() for row in sheet.iter_rows(min_row=header_no + 1, max_row=header_no + 2, values_only=True) for v in row)
                    stock = "нач. остаток" in subheaders or "начальный остаток" in subheaders
                    kind = "monthly_stock" if stock else "monthly_sales"
                    base = [fields.get("name"), fields.get("unit"), fields["code"]] if stock else [fields.get("name"), fields["code"]]
                    names = ["Номенклатура", "Ед.", "Номенклатура.Код"] if stock else ["Номенклатура", "Номенклатура.Код"]
                    views = [(kind, TableView(sheet, header_no, names + [raw[i] for i in month_indexes], base + month_indexes))]
                elif any("поступление до" in str(v or "").casefold() for v in raw):
                    arrival_indexes = [i for i, v in enumerate(raw) if "поступление до" in str(v or "").casefold()]
                    views = [("inbound", TableView(sheet, header_no, ["Код 1с", "Артикул ИЭК", "Наименование"] + [raw[i] for i in arrival_indexes],
                                                   [fields["code"], fields.get("article"), fields.get("name")] + arrival_indexes))]
                else:
                    views = flat_views(sheet, header_info)
                    found = {kind for kind, _ in views}
                    for field, expected in [("stock", "current_stock"), ("quantity", "sales_detail"), ("arrival_qty", "inbound_rows")]:
                        if field in fields and expected not in found:
                            issues.append(_warning("INCOMPLETE_SOURCE_COLUMNS", "error", f"{filename}:{sheet.title}: колонка {raw[fields[field]]} найдена, но не хватает однозначных даты/склада/полей операции для её использования."))
            elif _detect(filename, sheet) == "seasonality":
                views = [("seasonality", sheet)]
            if not views:
                kinds.append("unknown")
                unsupported.add(f"{filename}:{sheet.title}")
                issues.append(_warning("UNKNOWN_SHEET", "error", f"{filename}:{sheet.title}: структура не поддерживается; лист не включён в расчёт."))
                continue
            for kind, view in views:
                kinds.append("inbound" if kind == "inbound_rows" else kind)
                _parse_sheet(filename, view, kind, warehouse_scope, data, issues, counts)
        file_info.append({"file_name": filename, "sheet_names": workbook.sheetnames, "detected_type": kinds[0] if len(set(kinds)) == 1 else "mixed"})
        workbook.close()
    _merge_sources(data, issues, counts)
    sales_codes = {x["sku"] for x in data["sales_detail"]}
    for source_kind in ("monthly_sales", "monthly_stock", "inbound", "moq"):
        unmatched = {x["sku"] for x in data[source_kind]} - sales_codes
        if unmatched:
            issues.append(_warning("SOURCE_CODE_UNMATCHED", "warning", f"{source_kind}: {len(unmatched)} кодов 1С отсутствуют в детализации продаж; сопоставление по наименованию не выполняется."))
            counts[f"unmatched_codes_{source_kind}"] = len(unmatched)
    units_by_code = {x["sku"]: x["unit"] for x in data["sales_detail"] if x["unit"]}
    unit_mismatches = {x["sku"] for x in data["monthly_stock"] if x["sku"] in units_by_code and x["unit"] and x["unit"] != units_by_code[x["sku"]]}
    if unit_mismatches:
        issues.append(_warning("UNIT_MISMATCH", "warning", f"У {len(unit_mismatches)} кодов единицы в продажах и месячных остатках различаются."))
        counts["unit_mismatch_codes"] = len(unit_mismatches)
    detail_monthly: dict[tuple[str, str], float] = defaultdict(float)
    for sale in data["sales_detail"]:
        detail_monthly[(sale["sku"], sale["date"][:7])] += sale["quantity"]
    monthly_values: dict[tuple[str, str], float] = defaultdict(float)
    for sale in data["monthly_sales"]:
        monthly_values[(sale["sku"], sale["month"])] += sale["quantity"]
    compared = set(detail_monthly) & set(monthly_values)
    different = [key for key in compared if abs(detail_monthly[key] - monthly_values[key]) > max(1, abs(monthly_values[key]) * 0.01)]
    if different:
        issues.append(_warning("MONTHLY_DETAIL_MISMATCH", "warning", f"{len(different)} из {len(compared)} сопоставимых артикулов-месяцев расходятся более чем на 1%; месячный источник не суммируется с детализацией. Возможны различия охвата складов и возвратов."))
        counts["monthly_detail_mismatch"] = len(different)
    inbound_keys: dict[tuple, list[dict]] = defaultdict(list)
    for arrival in data["inbound"]:
        inbound_keys[(arrival["sku"], arrival["document"], arrival["expected_date"], arrival["quantity"])].append(arrival)
    ambiguous_inbound = 0
    for duplicate_rows in inbound_keys.values():
        if len(duplicate_rows) > 1:
            ambiguous_inbound += len(duplicate_rows)
            for arrival in duplicate_rows:
                arrival["duplicate_ambiguous"] = True
    if ambiguous_inbound:
        issues.append(_warning("AMBIGUOUS_INBOUND_DUPLICATES", "warning", f"{ambiguous_inbound} строк поставок имеют одинаковые код, документ, дату и количество; они исключаются из подтверждённого поступления до проверки."))
        counts["ambiguous_inbound_rows"] = ambiguous_inbound
    moq_by_code: dict[str, set[float]] = defaultdict(set)
    for condition in data["moq"]:
        moq_by_code[condition["sku"]].add(condition["minimum"])
    conflicting_moq = sum(len(values) > 1 for values in moq_by_code.values())
    if conflicting_moq:
        issues.append(_warning("CONFLICTING_MOQ", "warning", f"У {conflicting_moq} кодов есть разные значения MOQ; рекомендации по ним блокируются."))
        counts["conflicting_moq_codes"] = conflicting_moq
    if not data["current_stock"]:
        issues.append(_warning("CURRENT_STOCK_MISSING", "warning", "Текущий остаток не предоставлен. Прогноз спроса доступен; фактический заказ без остатка не определяется."))
        unsupported.add("current_available_stock")
    if not data["stockouts"]:
        issues.append(_warning("STOCKOUT_INTERVALS_MISSING", "warning", "Нет подтверждённых интервалов отсутствия; компенсация stockout не применяется."))
        unsupported.add("confirmed_stockout_intervals")
    if not data["categories"]:
        issues.append(_warning("CATEGORIES_MISSING", "warning", "Нет справочника категорий; применяются общие настройки сервиса."))
        unsupported.add("category_mapping")
    if not any(x.get("customer_id") for x in data["sales_detail"]):
        issues.append(_warning("CUSTOMER_ID_MISSING", "info", "В выгрузке продаж нет обезличенного ID клиента; проверка концентрации клиента недоступна."))
    issues.append(_warning("LEAD_TIME_DIRECTORY_MISSING", "warning", "Справочник сроков поставки не загружен; нужен явно заданный срок в запросе расчёта."))
    if data["seasonality"]:
        issues.append(_warning("SEASONALITY_AGGREGATE_REVENUE", "warning", "Коэффициенты сезонности в файле рассчитаны по агрегированной выручке, не по количеству каждого артикула."))
    # Keep the full normalized payload locally; cap response warnings, retaining
    # a summary count so the UI is usable even for a very large workbook.
    payload = {"data": data, "issues": issues, "row_counts": dict(counts), "files": file_info, "warehouse_scope": warehouse_scope}
    preview_id = save_preview(payload)
    visible = issues[:200]
    if len(issues) > 200:
        visible.append(_warning("ADDITIONAL_ISSUES", "warning", f"Ещё {len(issues) - 200} проблемных строк сохранено в локальном наборе."))
    return {"preview_id": preview_id, "files": file_info, "row_counts": dict(counts), "warnings": visible, "unsupported_fields": sorted(unsupported)}


def _future_seasonality(profile: dict[int, float], calculation_date: date, horizon: int) -> float:
    if not profile:
        return 1.0
    return sum(profile[(calculation_date + timedelta(days=day)).month] for day in range(1, horizon + 1)) / horizon


def dataset_to_inputs(dataset: dict, request: dict) -> tuple[list[CalculationInput], list[dict]]:
    data = dataset["data"]
    calculation_date = date.fromisoformat(request["calculation_date"])
    horizon = request["review_period_days"] + request["default_lead_time_days"]
    start = calculation_date - timedelta(days=89)
    profile = {int(month): content["coefficient"] for month, content in data["seasonality"].items()}
    if profile and len(profile) != 12:
        profile = {}
    by_sku: dict[str, list[dict]] = defaultdict(list)
    identity: dict[str, tuple[str, str]] = {}
    for record in data["sales_detail"]:
        sale_date = date.fromisoformat(record["date"])
        if start <= sale_date <= calculation_date:
            by_sku[record["sku"]].append(record)
            identity[record["sku"]] = (record["name"], record["unit"])
    stocks: dict[str, list[dict]] = defaultdict(list)
    for record in data["current_stock"]:
        stocks[record["sku"]].append(record)
    moqs: dict[str, list[dict]] = defaultdict(list)
    for record in data["moq"]:
        moqs[record["sku"]].append(record)
    arrivals: dict[str, list[dict]] = defaultdict(list)
    for record in data["inbound"]:
        arrivals[record["sku"]].append(record)
    categories = {record["sku"]: record["category"] for record in data["categories"]}
    stockouts: dict[str, list[dict]] = defaultdict(list)
    for record in data["stockouts"]:
        stockouts[record["sku"]].append(record)
    category_settings = {setting["category"]: setting for setting in request.get("category_settings", [])}
    warnings = list(dataset["issues"])
    blocked_skus = {w["sku"] for w in warnings if w.get("sku") and w["severity"] == "error"}
    inputs: list[CalculationInput] = []
    missing_stock_count = 0
    estimated_stock_count = 0
    monthly_stocks = defaultdict(list)
    for record in data["monthly_stock"]:
        if record["month"] == calculation_date.strftime("%Y-%m"):
            monthly_stocks[record["sku"]].append(record)
    for sku, records in by_sku.items():
        if sku in blocked_skus:
            continue
        stock_records = [x for x in stocks.get(sku, []) if date.fromisoformat(x["date"]) <= calculation_date]
        if not stock_records:
            missing_stock_count += 1
        latest_date = max((x["date"] for x in stock_records), default=None)
        latest = [x for x in stock_records if x["date"] == latest_date]
        if len(latest) > 1:
            warnings.append(_warning("AMBIGUOUS_CURRENT_STOCK", "error", "Несколько остатков с одной датой; требуется выбор склада/строки.", sku))
            continue
        name, unit = identity[sku]
        stock_estimate = None
        if not latest and request.get("stock_mode") == "monthly_estimate":
            openings = monthly_stocks.get(sku, [])
            if len(openings) == 1 and openings[0]["quantity"] >= 0 and normalize(openings[0].get("unit")) == normalize(unit):
                opening = openings[0]
                opening_date = date(calculation_date.year, calculation_date.month, 1)
                sold = sum(x["quantity"] for x in records if opening_date <= date.fromisoformat(x["date"]) <= calculation_date)
                estimated = max(0, opening["quantity"] - sold)
                stock_estimate = {"quantity": estimated, "opening_quantity": opening["quantity"], "deducted_sales": sold,
                                  "date": opening_date.isoformat(), "source": opening["source"], "method": "opening_minus_sales"}
                latest = [{"quantity": estimated}]
                estimated_stock_count += 1
                missing_stock_count -= 1
        category = categories.get(sku)
        setting = category_settings.get(category, {})
        moq_values = {x["minimum"] for x in moqs.get(sku, [])}
        if len(moq_values) > 1:
            warnings.append(_warning("CONFLICTING_MOQ", "error", "Для артикула несколько разных минимальных партий.", sku))
            continue
        stockout_intervals = tuple(Stockout(date.fromisoformat(x["start"]), date.fromisoformat(x["end"]), x["confirmed"]) for x in stockouts.get(sku, []))
        operations = tuple(Operation(quantity=x["quantity"], date=date.fromisoformat(x["date"]), customer_id=x.get("customer_id"), reference=f"{x['source']['file']}:{x['source']['sheet']}:{x['source']['row']}") for x in records)
        inbound_records = []
        for arrival in arrivals.get(sku, []):
            if arrival.get("duplicate_ambiguous"):
                warnings.append(_warning("AMBIGUOUS_INBOUND_FOR_SKU", "error", "Повторяющиеся строки поставки не учтены до проверки документа.", sku))
                continue
            if arrival["unit_conversion_uncertain"] and unit == "м":
                warnings.append(_warning("INBOUND_UNIT_CONVERSION_UNKNOWN", "error", "Поставка указана бухтами, учёт ведётся в метрах; коэффициент перевода не подтверждён.", sku))
                continue
            inbound_records.append(Inbound(arrival["quantity"], date.fromisoformat(arrival["expected_date"])))
        future_factor = _future_seasonality(profile, calculation_date, horizon)
        inputs.append(CalculationInput(
            sku=sku, item_name=name, supplier_id="iek", supplier_name="ИЭК", warehouse=request["warehouse_scope"], unit=unit,
            category=category, calculation_date=calculation_date, review_period_days=request["review_period_days"],
            lead_time_days=request["default_lead_time_days"], service_level_z=setting.get("service_level_z", request["service_level_z"]),
            available_stock=latest[0]["quantity"] if latest else None, available_stock_date=date.fromisoformat(latest_date) if latest_date else None,
            supplier_article=moqs[sku][0]["supplier_article"] if moqs.get(sku) else None,
            moq=next(iter(moq_values)) if moq_values else None, pack_multiple=None,
            growth_override_percent=setting.get("growth_override_percent", request.get("growth_override_percent", 0)),
            seasonality_coefficient=future_factor,
            seasonality_profile=tuple(profile[month] for month in range(1, 13)) if profile else None,
            operations=operations, inbound=tuple(inbound_records), stockouts=stockout_intervals,
            demand_window_days=90,
            stock_estimate=stock_estimate,
        ))
    if estimated_stock_count:
        warnings.append(_warning("MONTHLY_STOCK_ESTIMATE", "warning", f"Для {estimated_stock_count} позиций рассчитан предварительный план по оценке остатка на основе начала месяца и продаж. Приходы, возвраты, резервы и перемещения не известны; утверждение фактического заказа недоступно."))
    if missing_stock_count:
        warnings.append(_warning(
            "CURRENT_STOCK_MISSING_SUMMARY",
            "warning",
            f"Для {missing_stock_count} артикулов рассчитан прогноз спроса без текущего остатка; количество фактического заказа неизвестно.",
        ))
    if not profile:
        warnings.append(_warning("SEASONALITY_MISSING", "warning", "Коэффициенты сезонности не загружены; используется 1.0."))
    if not inputs:
        warnings.append(_warning("NO_CALCULABLE_ITEMS", "error", "Нет подходящей истории продаж на выбранную дату или все позиции имеют конфликтующие данные."))
    return inputs, warnings
