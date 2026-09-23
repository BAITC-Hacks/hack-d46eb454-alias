from openpyxl import Workbook

from backend.importer import _code, _date, _detect, _month, _same_warehouse, dataset_to_inputs
from backend import storage
from backend.replenishment import calculate_item
from fastapi.testclient import TestClient
from backend.main import app
from io import BytesIO
import pytest


def excel_bytes(tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in tables.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


SALES = [["Наименование товара", "Кол-во", "Код товара", "Тип документа", "Дата операции", "Ед. изм.", "Наименование склада", "Номер документа"],
         ["Синтетический товар", 90, "001_", "Расходная накладная", "2026-09-22", "шт", "Алматы", "TEST-1"],
         ["Синтетический товар", 90, "001_", "Расходная накладная", "2026-09-21", "шт", "Алматы", "TEST-2"]]
STOCK = [["Дата среза", "Код 1С", "Доступный остаток", "Склад", "Единица измерения"],
         ["2026-09-23", "001_", 10, "Алматы", "шт"]]
POLICY = [["MOQ", "Категория товара", "Код товара", "Артикул поставщика"], [5, "Свет", "001_", "IEK-TEST"]]
ARRIVAL = [["Код товара", "Количество в пути", "Дата поставки", "Документ поставки", "Ед."],
           ["001_", 20, "2026-09-30", "IN-1", "шт"]]


def import_dataset(client, tables):
    files = [("files", (name, excel_bytes(sheets), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")) for name, sheets in tables.items()]
    response = client.post("/api/import/preview", files=files, data={"warehouse_scope": "Алматы"})
    assert response.status_code == 200, response.text
    preview = response.json()
    commit = client.post("/api/import/commit", json={"preview_id": preview["preview_id"]})
    assert commit.status_code == 201, commit.text
    dataset = storage.get_dataset(commit.json()["dataset_id"])
    run = client.post("/api/recommendations/calculate", json={"dataset_id": commit.json()["dataset_id"], "calculation_date": "2026-09-23", "warehouse_scope": "Алматы", "review_period_days": 7, "default_lead_time_days": 21, "service_level_z": 0})
    assert run.status_code == 200, run.text
    return preview, dataset, run.json()


def test_split_files_multisheet_and_combined_table_produce_same_result(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    client = TestClient(app)
    sheets = {"Продажи": SALES, "Остатки": STOCK, "Условия": POLICY, "Поступления": ARRIVAL}
    split, data, split_run = import_dataset(client, {f"{name}.xlsx": {name: rows} for name, rows in sheets.items()})
    # The filename must not classify every sheet as MOQ.
    book, _, book_run = import_dataset(client, {"MOQ общий.xlsx": sheets})
    combined = [SALES[0] + ["Текущий остаток", "Дата остатка", "MOQ", "Категория", "Артикул поставщика"]]
    combined += [row + [10, "2026-09-23", 5, "Свет", "IEK-TEST"] for row in SALES[1:]]
    one, one_data, one_run = import_dataset(client, {"общая.xlsx": {"Все данные": combined, "Поступления": ARRIVAL}})
    assert not any(w["severity"] == "error" for w in split["warnings"]), split
    assert {f["detected_type"] for f in book["files"]} == {"mixed"}
    assert data["data"]["sales_detail"][0]["sku"] == "001_"
    assert len(one_data["data"]["current_stock"]) == 1
    assert len(one_data["data"]["categories"]) == 1
    assert one["row_counts"]["sales_detail"] == 2
    items = [r["suppliers"][0]["items"][0] for r in (split_run, book_run, one_run)]
    assert items[0]["components"] == items[1]["components"] == items[2]["components"]
    assert items[0]["components"]["available_stock"] == 10
    assert items[0]["components"]["eligible_inbound"][0]["quantity"] == 20
    assert items[0]["category"] == "Свет"


def test_combined_one_sheet_and_reordered_headers_with_title(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    combined = [["Отчёт закупщика"], [], SALES[0] + ["Остаток", "Дата остатка", "MOQ", "Категория"]]
    combined += [row + [10, "2026-09-23", 5, "Свет"] for row in SALES[1:]]
    preview, data, run = import_dataset(TestClient(app), {"произвольное имя.xlsx": {"Всё": combined}})
    assert run["totals"]["item_count"] == 1
    assert data["data"]["sales_detail"][0]["source"]["row"] == 4
    assert preview["row_counts"]["current_stock"] == 1


@pytest.mark.parametrize("change,code", [("stock", "SOURCE_VALUE_CONFLICT"), ("unit", "SOURCE_UNIT_CONFLICT"), ("category", "SOURCE_VALUE_CONFLICT"), ("overlap", "OVERLAPPING_SALES_SOURCES")])
def test_conflicting_sources_block_approval(tmp_path, monkeypatch, change, code):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    sheets = {"Продажи": SALES, "Остатки": STOCK, "Условия": POLICY}
    if change == "stock":
        sheets["Другой остаток"] = [STOCK[0], ["2026-09-23", "001_", 999, "Алматы", "шт"]]
    elif change == "unit":
        sheets["Остатки"] = [STOCK[0], ["2026-09-23", "001_", 10, "Алматы", "м"]]
    elif change == "category":
        sheets["Категории"] = [["Код", "Категория"], ["001_", "Другая"]]
    else:
        sheets["Продажи повторно"] = SALES
    client = TestClient(app)
    preview, _, run = import_dataset(client, {"файл.xlsx": sheets})
    assert any(w["code"] == code for w in preview["warnings"])
    assert run["totals"]["item_count"] == 0
    assert client.post(f"/api/orders/{run['run_id']}/approve", json={"approved_by": "test"}).status_code == 409


def test_unknown_and_ambiguous_columns_not_guessed(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    preview, _, run = import_dataset(TestClient(app), {"moq.xlsx": {"Странная таблица": [["Имя", "Число"], ["товар", 99]],
                                                                            "Дубли": [["Код", "Код товара", "MOQ"], ["1", "2", 5]]}})
    assert {"UNKNOWN_SHEET", "AMBIGUOUS_COLUMNS"} <= {w["code"] for w in preview["warnings"]}
    assert run["totals"]["item_count"] == 0


def test_same_sale_requisites_are_not_silently_dropped(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    preview, dataset, run = import_dataset(TestClient(app), {"повторы.xlsx": {"Продажи": [SALES[0], SALES[1], SALES[1]], "Остатки": STOCK}})
    assert len(dataset["data"]["sales_detail"]) == 2
    assert any(w["code"] == "DUPLICATE_SALE_ROW" and w["severity"] == "error" for w in preview["warnings"])
    assert run["totals"]["item_count"] == 0


def test_malformed_seasonality_reports_error_instead_of_crashing(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    preview, _, _ = import_dataset(TestClient(app), {"сезонность.xlsx": {"Сезонность": [["Месяц"], ["Январь"], ["Норм. коэф."], [1.0]]}})
    assert any(w["code"] == "INVALID_SEASONALITY" for w in preview["warnings"])


def test_1c_code_and_date_are_not_normalized_away():
    assert _code("0012300_") == "0012300_"
    assert _date("22.09.2026 16:15:38").isoformat() == "2026-09-22"
    assert _month("сент. 2026") == "2026-09"


def test_warehouse_match_ignores_case_and_outer_whitespace():
    assert _same_warehouse("Алматы", "алматы")
    assert _same_warehouse("  АЛМАТЫ  ", "Алматы")
    assert not _same_warehouse("Астана", "Алматы")


def test_detects_monthly_sources_from_structure_when_filename_is_mojibake():
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Лист_1"
    sales.append(["Номенклатура", "Номенклатура.Код", "янв. 2024", "февр. 2024", "март 2024", "апр. 2024", "май 2024", "июнь 2024"])
    assert _detect("broken-name.xlsx", sales) == "monthly_sales"

    stock = workbook.create_sheet("Остатки")
    stock.append(["Номенклатура", "Ед.", "Номенклатура.Код", "янв. 2024", "февр. 2024", "март 2024", "апр. 2024", "май 2024", "июнь 2024"])
    assert _detect("broken-name.xlsx", stock) == "monthly_stock"


def test_detects_seasonality_from_sheet_structure_when_filename_is_mojibake():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сезонность"
    for _ in range(9):
        sheet.append([])
    sheet.append([None, "Месяц", "Продажи 2024", "Коэф. сезонности"])
    for _ in range(16):
        sheet.append([])
    sheet.append([None, "Месяц", "Коэф. 2025", "Коэф. 2026", "Средний коэф.", "Норм. коэф."])
    assert _detect("broken-name.xlsx", sheet) == "seasonality"


def test_dataset_forecasts_without_current_stock():
    dataset = {
        "data": {
            "sales_detail": [{"sku": "001_", "date": "2026-09-22", "quantity": 10, "name": "Тест", "unit": "шт", "customer_id": None,
                              "source": {"file": "synthetic.xlsx", "sheet": "Продажи", "row": 2}}],
            "monthly_sales": [], "monthly_stock": [], "inbound": [], "moq": [], "seasonality": {}, "current_stock": [], "categories": [], "stockouts": [],
        },
        "issues": [],
    }
    request = {"calculation_date": "2026-09-23", "warehouse_scope": "Алматы", "review_period_days": 7,
               "default_lead_time_days": 21, "service_level_z": 1.65}
    inputs, warnings = dataset_to_inputs(dataset, request)
    assert len(inputs) == 1
    assert calculate_item(inputs[0])["recommended_quantity"] is None
    assert any(w["code"] == "CURRENT_STOCK_MISSING_SUMMARY" for w in warnings)
    dataset["data"]["current_stock"] = [{"sku": "001_", "quantity": 5, "date": "2026-09-22"}]
    inputs, warnings = dataset_to_inputs(dataset, request)
    assert len(inputs) == 1
    assert inputs[0].sku == "001_"
    assert inputs[0].available_stock == 5
    baseline = calculate_item(inputs[0])
    dataset["data"]["categories"] = [{"sku": "001_", "category": "Критичный"}]
    request["category_settings"] = [{"category": "Критичный", "service_level_z": 2.5, "growth_override_percent": 10}]
    categorized, _ = dataset_to_inputs(dataset, request)
    assert categorized[0].category == "Критичный"
    assert categorized[0].service_level_z == 2.5
    assert categorized[0].growth_override_percent == 10
    changed = calculate_item(categorized[0])
    assert changed["components"]["safety_stock"] > baseline["components"]["safety_stock"]
    assert changed["components"]["forecast_over_horizon"] > baseline["components"]["forecast_over_horizon"]


def test_monthly_estimate_is_opt_in_and_cannot_be_approved_as_actual_order(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    client = TestClient(app)
    monthly = [["Номенклатура", "Ед.", "Код", "сент. 2026"], [None, None, None, "Количество"],
               [None, None, None, "Нач. остаток"], ["Тест", "шт", "001_", 200]]
    _, dataset, forecast = import_dataset(client, {"синтетический.xlsx": {"Продажи": SALES, "Месячные остатки": monthly}})
    item = forecast["suppliers"][0]["items"][0]
    assert item["recommended_quantity"] is None
    assert item["components"]["forecast_over_horizon"] == 56
    assert client.patch(f"/api/recommendations/runs/{forecast['run_id']}/items/{item['item_id']}/adjust", json={"approved_quantity": 1, "reason": "test", "changed_by": "test"}).status_code == 409
    request = {"dataset_id": forecast["dataset_id"], "calculation_date": "2026-09-23", "warehouse_scope": "Алматы", "review_period_days": 7, "default_lead_time_days": 21, "service_level_z": 0, "stock_mode": "monthly_estimate"}
    response = client.post("/api/recommendations/calculate", json=request)
    assert response.status_code == 200
    run = response.json()
    item = run["suppliers"][0]["items"][0]
    assert item["components"]["stock_estimate"]["quantity"] == 20  # 200 - (90 + 90)
    assert item["components"]["available_stock"] is None  # Never presented as fact.
    assert item["recommended_quantity"] == 36  # 56 forecast - 20 estimate
    assert "estimated_stock" in item["flags"]
    assert client.post(f"/api/orders/{run['run_id']}/approve", json={"approved_by": "test"}).status_code == 409
    assert client.get(f"/api/orders/{run['run_id']}/export?format=csv").status_code == 409
    for month in ("2026-08", "2026-10"):
        dataset["data"]["monthly_stock"][0]["month"] = month
        inputs, _ = dataset_to_inputs(dataset, request)
        assert inputs[0].available_stock is None  # Neither stale nor future stock is substituted.
    dataset["data"]["monthly_stock"][0]["month"] = "2026-09"
    dataset["data"]["current_stock"] = [{"sku": "001_", "quantity": 10, "date": "2026-09-23"}]
    inputs, _ = dataset_to_inputs(dataset, request)
    assert inputs[0].available_stock == 10
    assert inputs[0].stock_estimate is None  # Actual snapshot takes precedence.


def test_preview_commit_is_idempotent(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    preview_id = storage.save_preview({"data": {"sales_detail": []}, "issues": []})
    first = storage.commit_preview(preview_id, "synthetic")
    second = storage.commit_preview(preview_id, "different label")
    assert second == first
